from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from rest_framework.test import APITestCase

from clients.models import Client
from products.models import Category, PriceLevel, PriceTier, Product, ProductPrice
from users.models import Role
from .models import Quotation
from .serializers import QuotationSerializer
from .transfer import HEADERS


class QuotationTransferTests(APITestCase):
    def setUp(self):
        self.role = Role.objects.create(name="Importador", code="IMPORT")
        self.grant("view_quotation", "add_quotation", "import_quotation", "export_quotation", "view_client", "view_product")
        self.user = get_user_model().objects.create_user(username="importer", email="importer@example.com", role=self.role)
        self.client.force_authenticate(self.user)
        level = PriceLevel.objects.create(name="Normal", code="NORMAL")
        self.customer = Client.objects.create(name="Cliente", tax_id="00123", price_level=level, department="La Paz", city_zone="Centro", whatsapp="70000000", client_type="Empresa", business_activity="Comercio")
        category = Category.objects.create(name="Equipos", code="EQ")
        self.product = Product.objects.create(sku="EQ-1", name="Equipo", category=category, normal_unit_price=20)

    def grant(self, *codes):
        self.role.permissions.add(*Permission.objects.filter(codename__in=codes))

    def row(self, reference="REC-1", serial="S001", **changes):
        row = dict(zip(HEADERS, [reference, "00123", "Cliente", "21/09/2026", "EQ-1", "Equipo", serial, "", "Pendiente", "Entrega"] ))
        row.update(changes)
        return row

    def upload(self, rows, mode="partial"):
        workbook = openpyxl.Workbook()
        workbook.active.append(HEADERS)
        for row in rows:
            workbook.active.append([row.get(header, "") for header in HEADERS])
        buffer = BytesIO()
        workbook.save(buffer)
        return self.client.post("/api/v1/quotations/import/", {"file": SimpleUploadedFile("cotizaciones.xlsx", buffer.getvalue()), "mode": mode}, format="multipart")

    def test_groups_units_uses_pricing_and_ignores_blank_rows(self):
        tier = PriceTier.objects.create(price_level=self.customer.price_level, minimum_quantity=2)
        ProductPrice.objects.create(product=self.product, price_tier=tier, price_level=self.customer.price_level, minimum_quantity=2, unit_price=15)
        result = self.upload([self.row(), {}, self.row(serial="S002")])
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.data["created"], 1)
        quotation = Quotation.objects.get()
        self.assertEqual(quotation.total, Decimal("30"))
        self.assertEqual(quotation.items.count(), 2)
        self.assertEqual(quotation.created_by, self.user)
        self.assertEqual(quotation.quotation_date.isoformat(), "2026-09-21")
        self.assertTrue(quotation.template_snapshot)
        self.assertTrue(quotation.number.startswith("COT-"))

    def test_partial_rejects_whole_invalid_group_and_preserves_row_numbers(self):
        result = self.upload([self.row(), {}, self.row(SKU="missing"), self.row("REC-2")])
        self.assertEqual(result.data["created"], 1)
        self.assertEqual(result.data["rejected"], 1)
        self.assertEqual(Quotation.objects.get().import_reference, "rec-2")
        self.assertTrue(any(item["fila"] == 4 for item in result.data["errors"]))

    def test_total_does_not_save_valid_groups_when_another_fails(self):
        result = self.upload([self.row(), self.row("REC-2", SKU="missing")], "total")
        self.assertEqual(result.data["created"], 0)
        self.assertEqual(result.data["rejected"], 2)
        self.assertFalse(Quotation.objects.exists())

    def test_repeated_reference_is_rejected_case_insensitively(self):
        self.upload([self.row()])
        result = self.upload([self.row("rec-1")])
        self.assertEqual(result.data["created"], 0)
        self.assertEqual(Quotation.objects.count(), 1)

    def test_invalid_fields_and_inconsistent_group_are_rejected(self):
        for changes in ({"Nombre_cliente": ""}, {"Nombre_cliente": "   "}, {"Fecha_cotizacion": ""}, {"NIT/CI": "missing"}, {"Nombre_cliente": "Otro"}, {"SKU": ""}, {"Numero_serie": "x" * 121}, {"Estado": "otro"}):
            with self.subTest(changes=changes):
                result = self.upload([self.row(**changes)])
                self.assertEqual(result.data["created"], 0)
                self.assertTrue(result.data["errors"])
        result = self.upload([self.row(), self.row(serial="S002", Observaciones="Distinta")])
        self.assertEqual(result.data["created"], 0)
        result = self.upload([self.row(), self.row()])
        self.assertEqual(result.data["created"], 0)

    def test_price_and_paid_state_require_separate_permissions(self):
        result = self.upload([self.row(Precio_unitario_Bs="0")])
        self.assertEqual(result.data["created"], 0)
        result = self.upload([self.row(Estado="Pagada")])
        self.assertEqual(result.data["created"], 0)
        self.grant("change_quotation", "change_quotation_item_price")
        result = self.upload([self.row(Precio_unitario_Bs="0", Estado="Pagada")])
        self.assertEqual(result.data["created"], 1, result.data)
        self.assertEqual(Quotation.objects.get().total, Decimal("0"))
        self.assertEqual(Quotation.objects.get().status, "paid")
        for price in ("NaN", "-1", "1.001", "10000000000"):
            result = self.upload([self.row("REC-2", Precio_unitario_Bs=price)])
            self.assertEqual(result.data["created"], 0)

    def test_import_export_permissions_are_independent(self):
        self.role.permissions.remove(Permission.objects.get(codename="export_quotation"))
        self.assertEqual(self.client.get("/api/v1/quotations/export/").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/quotations/template/").status_code, 200)
        self.grant("export_quotation")
        self.role.permissions.remove(Permission.objects.get(codename="import_quotation"))
        self.assertEqual(self.upload([self.row()]).status_code, 403)
        self.assertEqual(self.client.post("/api/v1/quotations/report/", {"errors": []}, format="json").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/quotations/export/").status_code, 200)

    def test_missing_permission_dependency_denies_import(self):
        self.role.permissions.remove(Permission.objects.get(codename="add_quotation"))
        self.assertEqual(self.upload([self.row()]).status_code, 403)

    def test_export_filters_and_reimport_does_not_duplicate(self):
        self.upload([self.row(Observaciones="=1+1"), self.row("REC-2")])
        quotation = Quotation.objects.get(import_reference="rec-1")
        quotation.notes = "=1+1"
        quotation.save()
        result = self.client.get("/api/v1/quotations/export/", {"search": quotation.number})
        self.assertEqual(result.status_code, 200)
        workbook = openpyxl.load_workbook(BytesIO(result.content))
        rows = list(workbook.active.values)
        self.assertEqual(len(rows), 2)
        self.assertEqual(list(rows[0]), HEADERS)
        self.assertEqual(workbook.active.cell(2, 10).data_type, "s")
        self.assertEqual(rows[1][1], "00123")
        self.assertEqual(rows[1][7], 20)
        result = self.upload([dict(zip(HEADERS, rows[1]))])
        self.assertEqual(result.data["created"], 0)
        self.assertEqual(Quotation.objects.count(), 2)

    def test_template_empty_malformed_report_and_native_date(self):
        template = self.client.get("/api/v1/quotations/template/")
        workbook = openpyxl.load_workbook(BytesIO(template.content))
        self.assertEqual(list(workbook.active.values), [tuple(HEADERS)])
        self.assertEqual(self.upload([{}]).status_code, 400)
        self.assertEqual(self.upload([self.row()], "bad").status_code, 400)
        result = self.upload([self.row(Fecha_cotizacion=datetime(2026, 9, 21))])
        self.assertEqual(result.data["created"], 1)
        self.assertEqual(self.client.post("/api/v1/quotations/report/", {"errors": [1]}, format="json").status_code, 400)
        self.assertEqual(self.client.post("/api/v1/quotations/report/", {"errors": [{"fila": 3, "motivo": "Prueba"}]}, format="json").status_code, 200)

    def test_total_rolls_back_a_late_database_conflict(self):
        original = QuotationSerializer.save
        calls = 0

        def save(serializer, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise IntegrityError("concurrent reference")
            return original(serializer, **kwargs)

        with patch.object(QuotationSerializer, "save", save):
            result = self.upload([self.row(), self.row("REC-2")], "total")
        self.assertEqual(result.status_code, 409)
        self.assertFalse(Quotation.objects.exists())

    def test_xls_native_date_and_corrupt_upload(self):
        import xlwt
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Cotizaciones")
        row = self.row(Fecha_cotizacion=datetime(2026, 9, 21))
        for index, header in enumerate(HEADERS):
            sheet.write(0, index, header)
            if header == "Fecha_cotizacion":
                sheet.write(1, index, row[header], xlwt.easyxf(num_format_str="DD/MM/YYYY"))
            else:
                sheet.write(1, index, row[header])
        buffer = BytesIO()
        workbook.save(buffer)
        response = self.client.post("/api/v1/quotations/import/", {"file": SimpleUploadedFile("cotizaciones.xls", buffer.getvalue())}, format="multipart")
        self.assertEqual(response.data["created"], 1)
        for name in ("bad.xlsx", "bad.csv"):
            response = self.client.post("/api/v1/quotations/import/", {"file": SimpleUploadedFile(name, b"invalid")}, format="multipart")
            self.assertEqual(response.status_code, 400)
