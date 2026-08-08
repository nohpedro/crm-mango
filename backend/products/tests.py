from django.core.files.uploadedfile import SimpleUploadedFile
from io import BytesIO
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework import serializers
from rest_framework.test import APITestCase
from rest_framework import status

from users.models import User

from .models import (
    Category,
    PriceLevel,
    PriceTier,
    Product,
    ProductImage,
    ProductPrice,
)
from .serializers import (
    PriceTierSerializer,
    ProductImageSerializer,
    ProductPriceSerializer,
)


class ProductImageValidationTests(TestCase):
    def test_rejects_unsupported_image_type(self):
        image = SimpleUploadedFile(
            "catalog.gif",
            b"gif-bytes",
            content_type="image/gif",
        )

        with self.assertRaises(serializers.ValidationError):
            ProductImageSerializer().validate_image(image)


class ProductPriceSerializerTests(TestCase):
    def test_derives_legacy_fields_from_the_selected_price_tier(self):
        category = Category.objects.create(name="General", code="GEN")
        product = Product.objects.create(category=category, name="Producto", sku="PROD-001")
        level = PriceLevel.objects.create(name="Mayorista", code="MAY")
        tier = PriceTier.objects.create(price_level=level, minimum_quantity=3)

        serializer = ProductPriceSerializer(
            data={
                "product": product.id,
                "price_tier": tier.id,
                "unit_price": "15.00",
                "discount_percent": "0",
                "is_active": True,
            }
        )

        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertEqual(serializer.validated_data["price_level"], level)
        self.assertEqual(serializer.validated_data["minimum_quantity"], 3)

    def test_price_tier_quantity_errors_are_in_spanish(self):
        level = PriceLevel.objects.create(name="Preferencial", code="PREF")
        PriceTier.objects.create(price_level=level, minimum_quantity=3)
        invalid = PriceTierSerializer(
            data={
                "price_level": level.id,
                "minimum_quantity": "tres",
                "is_active": True,
            }
        )
        duplicate = PriceTierSerializer(
            data={
                "price_level": level.id,
                "minimum_quantity": 3,
                "is_active": True,
            }
        )

        self.assertFalse(invalid.is_valid())
        self.assertIn(
            "número entero",
            str(invalid.errors["minimum_quantity"][0]),
        )
        self.assertFalse(duplicate.is_valid())
        self.assertIn(
            "Ya existe un nivel",
            str(duplicate.errors["minimum_quantity"][0]),
        )

    def test_editing_tier_updates_its_product_prices(self):
        category = Category.objects.create(name="Equipos", code="EQ")
        product = Product.objects.create(
            category=category,
            name="Producto",
            sku="EQ-001",
        )
        level = PriceLevel.objects.create(name="Mayorista", code="MAY")
        tier = PriceTier.objects.create(price_level=level, minimum_quantity=3)
        price = ProductPrice.objects.create(
            product=product,
            price_level=level,
            price_tier=tier,
            minimum_quantity=3,
            unit_price="15.00",
        )
        serializer = PriceTierSerializer(
            tier,
            data={"minimum_quantity": 5},
            partial=True,
        )

        self.assertTrue(serializer.is_valid(), serializer.errors)
        serializer.save()

        price.refresh_from_db()
        self.assertEqual(price.minimum_quantity, 5)


    def test_rejects_images_larger_than_five_megabytes(self):
        image = SimpleUploadedFile(
            "catalog.jpg",
            b"x" * (5 * 1024 * 1024 + 1),
            content_type="image/jpeg",
        )

        with self.assertRaises(serializers.ValidationError):
            ProductImageSerializer().validate_image(image)


class PriceTierDeletionAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="catalog-admin",
            email="catalog-admin@example.com",
            password="Demo12345!",
        )
        self.client.force_authenticate(self.user)
        category = Category.objects.create(name="General", code="GEN-DELETE")
        self.product = Product.objects.create(
            category=category,
            name="Producto con precio especial",
            sku="PRICE-TIER-DELETE",
        )
        self.level = PriceLevel.objects.create(name="Mayorista", code="MAY-DELETE")

    def test_deleting_a_tier_removes_its_associated_special_prices(self):
        tier = PriceTier.objects.create(price_level=self.level, minimum_quantity=3)
        price = ProductPrice.objects.create(
            product=self.product,
            price_level=self.level,
            price_tier=tier,
            minimum_quantity=3,
            unit_price="15.00",
        )

        response = self.client.delete(f"/api/v1/catalog/price-tiers/{tier.id}/")

        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(PriceTier.objects.filter(pk=tier.pk).exists())
        self.assertFalse(ProductPrice.objects.filter(pk=price.pk).exists())

    def test_cannot_delete_the_required_x1_tier(self):
        required_tier = PriceTier.objects.get(
            price_level=self.level,
            minimum_quantity=1,
        )

        response = self.client.delete(
            f"/api/v1/catalog/price-tiers/{required_tier.id}/"
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertTrue(PriceTier.objects.filter(pk=required_tier.pk).exists())


class ProductRestoreAPITests(APITestCase):
    def test_archived_product_can_be_restored(self):
        admin = User.objects.create_superuser(
            username="catalog-admin",
            email="catalog-admin@example.com",
            password="Password123!",
        )
        category = Category.objects.create(name="General", code="GEN")
        product = Product.objects.create(
            category=category,
            name="Producto archivado",
            sku="ARCH-001",
        )
        product.soft_delete()
        self.client.force_authenticate(admin)

        response = self.client.post(
            f"/api/v1/catalog/products/{product.pk}/restore/",
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        product.refresh_from_db()
        self.assertIsNone(product.deleted_at)
        self.assertTrue(product.is_active)

    def test_product_cover_can_be_selected_with_a_json_patch(self):
        admin = User.objects.create_superuser(
            username="image-admin",
            email="image-admin@example.com",
            password="Password123!",
        )
        category = Category.objects.create(name="Imágenes", code="IMG")
        product = Product.objects.create(
            category=category,
            name="Producto con galería",
            sku="IMG-001",
        )
        current = ProductImage.objects.create(
            product=product,
            image="products/current.jpg",
            is_primary=True,
        )
        replacement = ProductImage.objects.create(
            product=product,
            image="products/replacement.jpg",
            is_primary=False,
        )
        self.client.force_authenticate(admin)

        response = self.client.patch(
            f"/api/v1/catalog/product-images/{replacement.pk}/",
            {"is_primary": True},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        current.refresh_from_db()
        replacement.refresh_from_db()
        self.assertFalse(current.is_primary)
        self.assertTrue(replacement.is_primary)

    def test_product_import_rejects_duplicate_sku_in_same_file(self):
        admin = User.objects.create_superuser(
            username="product-import-admin",
            email="product-import-admin@example.com",
            password="Password123!",
        )
        category = Category.objects.create(name="General", code="GEN")
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["SKU", "Nombre", "Categoría", "Código de barras", "Descripción", "Precio de venta normal (nivel x1) (Bs)", "Activo", "Almacén", "Cantidad"])
        sheet.append(["SKU-001", "Producto uno", "GEN", "", "", "100.00", "SI", "", ""])
        sheet.append(["SKU-001", "Producto dos", "GEN", "", "", "150.50", "SI", "", ""])
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile("productos.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.client.force_authenticate(admin)
        response = self.client.post("/api/v1/catalog/import/", {"file": upload, "mode": "partial"}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["created"], 1)
        self.assertEqual(response.data["rejected"], 1)
        self.assertEqual(Product.objects.filter(category=category).count(), 1)


class ProductTransferTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            username="product-transfer-admin",
            email="product-transfer-admin@example.com",
            password="Password123!",
        )
        self.client.force_authenticate(self.admin)

    def test_import_creates_many_missing_categories_without_leaving_flow(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "SKU",
            "Nombre",
            "Categoría",
            "Código de barras",
            "Descripción",
            "Precio de venta normal (nivel x1) (Bs)",
            "Activo",
            "Almacén",
            "Cantidad",
        ])
        sheet.append(["AUTO-001", "Producto uno", "Calefacción", "", "", "100.00", "SI", "", ""])
        sheet.append(["AUTO-002", "Producto dos", "Herramientas", "", "", "250.50", "SI", "", ""])
        sheet.append(["AUTO-003", "Producto tres", "Calefacción", "", "", "75", "SI", "", ""])
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "productos.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/v1/catalog/import/",
            {
                "file": upload,
                "mode": "partial",
                "create_missing_categories": "true",
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["created"], 3)
        self.assertEqual(response.data["rejected"], 0)
        self.assertCountEqual(
            response.data["created_categories"], ["Calefacción", "Herramientas"]
        )
        self.assertEqual(Category.objects.count(), 2)
        self.assertEqual(Product.objects.count(), 3)
        self.assertEqual(
            str(Product.objects.get(sku="AUTO-002").normal_unit_price),
            "250.50",
        )

    def test_template_and_export_include_normal_sale_price(self):
        category = Category.objects.create(name="General", code="GENERAL")
        Product.objects.create(
            category=category,
            name="Producto con precio",
            sku="PRICE-001",
            normal_unit_price="123.45",
        )

        template_response = self.client.get("/api/v1/catalog/template/")
        export_response = self.client.get("/api/v1/catalog/export/")

        self.assertEqual(template_response.status_code, status.HTTP_200_OK)
        template = load_workbook(BytesIO(template_response.content), read_only=True)
        template_headers = [cell.value for cell in next(template.active.iter_rows())]
        self.assertIn("Precio de venta normal (nivel x1) (Bs)", template_headers)

        self.assertEqual(export_response.status_code, status.HTTP_200_OK)
        exported = load_workbook(BytesIO(export_response.content), read_only=True)
        rows = list(exported.active.iter_rows(values_only=True))
        price_index = rows[0].index("Precio de venta normal (nivel x1) (Bs)")
        self.assertEqual(rows[1][price_index], 123.45)

    def test_import_rejects_empty_or_invalid_normal_sale_price(self):
        category = Category.objects.create(name="General", code="GENERAL")
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "SKU",
            "Nombre",
            "Categoría",
            "Código de barras",
            "Descripción",
            "Precio de venta normal (nivel x1) (Bs)",
            "Activo",
            "Almacén",
            "Cantidad",
        ])
        sheet.append(["NO-PRICE", "Sin precio", category.code, "", "", "", "SI", "", ""])
        sheet.append(["BAD-PRICE", "Precio inválido", category.code, "", "", "abc", "SI", "", ""])
        output = BytesIO()
        workbook.save(output)
        upload = SimpleUploadedFile(
            "productos.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/v1/catalog/import/",
            {"file": upload, "mode": "partial"},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["created"], 0)
        self.assertEqual(response.data["rejected"], 2)
        self.assertTrue(
            all(
                error["columna"] == "Precio de venta normal (nivel x1) (Bs)"
                for error in response.data["errors"]
            )
        )
