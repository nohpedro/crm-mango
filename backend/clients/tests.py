from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from products.models import PriceLevel
from users.models import User

from .models import Client, ClientType


def excel_file(headers, rows, name="clientes.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(name, output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class ClientTransferTests(APITestCase):
    headers = [
        "Nombre o razón social", "NIT/CI", "Departamento", "Ciudad/Zona", "WhatsApp",
        "Tipo de cliente", "Nivel de precio", "Rubro o actividad", "Observaciones",
    ]

    def setUp(self):
        self.admin = User.objects.create_superuser(
            username="import-admin", email="import-admin@example.com", password="Password123!"
        )
        self.level = PriceLevel.objects.create(name="Regular", code="REGULAR")
        self.client.force_authenticate(self.admin)

    def test_partial_import_saves_valid_rows_and_reports_invalid_rows(self):
        upload = excel_file(self.headers, [
            ["Cliente válido", "900000001", "La Paz", "La Paz", "70000001", "Tienda o comercio", "REGULAR", "Comercio", ""],
            ["", "900000002", "La Paz", "La Paz", "70000002", "Tienda o comercio", "REGULAR", "Comercio", ""],
        ])
        response = self.client.post("/api/v1/clients/import/", {"file": upload, "mode": "partial"}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["created"], 1)
        self.assertEqual(response.data["rejected"], 1)
        self.assertEqual(Client.objects.count(), 1)
        self.assertEqual(response.data["errors"][0]["fila"], 3)

    def test_template_has_reference_headers(self):
        response = self.client.get("/api/v1/clients/template/?file_format=xlsx")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual([cell.value for cell in next(workbook.active.iter_rows())], self.headers)

    def test_template_is_xlsx_only(self):
        response = self.client.get("/api/v1/clients/template/?file_format=xls")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_import_can_create_many_missing_references_in_one_process(self):
        upload = excel_file(self.headers, [
            ["Cliente Norte", "900000011", "La Paz", "Norte", "70000011", "Ferretería", "Especial Norte", "Comercio", ""],
            ["Cliente Sur", "900000012", "Tarija", "Sur", "70000012", "Constructor", "Especial Sur", "Construcción", ""],
            ["Cliente Norte 2", "900000013", "La Paz", "Norte", "70000013", "Ferretería", "Especial Norte", "Comercio", ""],
        ])

        response = self.client.post(
            "/api/v1/clients/import/",
            {
                "file": upload,
                "mode": "partial",
                "create_missing_client_types": "true",
                "create_missing_price_levels": "true",
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["created"], 3)
        self.assertEqual(response.data["rejected"], 0)
        self.assertCountEqual(
            response.data["created_client_types"], ["Ferretería", "Constructor"]
        )
        self.assertCountEqual(
            response.data["created_price_levels"], ["Especial Norte", "Especial Sur"]
        )
        self.assertTrue(ClientType.objects.filter(name="Ferretería").exists())
        self.assertTrue(PriceLevel.objects.filter(name="Especial Norte").exists())
        self.assertEqual(Client.objects.count(), 3)

    def test_observation_report_uses_documented_client_endpoint(self):
        response = self.client.post(
            "/api/v1/clients/report/",
            {
                "errors": [
                    {
                        "fila": 2,
                        "columna": "Nivel de precio",
                        "valor": "No existe",
                        "motivo": "No existe un nivel activo.",
                    }
                ]
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
