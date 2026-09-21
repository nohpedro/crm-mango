from io import BytesIO


class ExcelDependencyError(RuntimeError):
    pass


def _load_openpyxl():
    try:
        import openpyxl
    except ImportError as error:
        raise ExcelDependencyError(
            "La importación/exportación Excel requiere instalar openpyxl, xlrd y xlwt."
        ) from error
    return openpyxl


def read_rows(uploaded_file, expected_headers, date_headers=()):
    filename = (uploaded_file.name or "").lower()
    if not filename.endswith((".xlsx", ".xls")):
        raise ValueError("Solo se permiten archivos .xls o .xlsx.")

    raw = uploaded_file.read()
    if filename.endswith(".xlsx"):
        openpyxl = _load_openpyxl()
        try:
            workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
        except Exception as error:
            raise ValueError("El archivo no es un libro Excel .xlsx válido.") from error
        values = [[_cell(value) for value in row] for row in rows]
    else:
        try:
            import xlrd
        except ImportError as error:
            raise ExcelDependencyError(
                "La importación .xls requiere instalar xlrd."
            ) from error
        try:
            workbook = xlrd.open_workbook(file_contents=raw)
            sheet = workbook.sheet_by_index(0)
            values = [[_cell(sheet.cell_value(row, col)) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
            for col in range(sheet.ncols):
                if values and values[0][col] in date_headers:
                    for row in range(1, sheet.nrows):
                        if sheet.cell_type(row, col) == xlrd.XL_CELL_DATE:
                            values[row][col] = xlrd.xldate_as_datetime(sheet.cell_value(row, col), workbook.datemode).date().isoformat()
        except Exception as error:
            raise ValueError("El archivo no es un libro Excel .xls válido.") from error

    if not values:
        raise ValueError("El archivo está vacío.")
    headers = [_normalize_header(value) for value in values[0]]
    expected = [_normalize_header(value) for value in expected_headers]
    if headers != expected:
        missing = [header for header in expected_headers if _normalize_header(header) not in headers]
        extra = [header for header in headers if header not in expected]
        details = []
        if missing:
            details.append(f"Faltan: {', '.join(missing)}")
        if extra:
            details.append(f"No reconocidas: {', '.join(extra)}")
        raise ValueError("Los encabezados no coinciden con la plantilla. " + " ".join(details))

    parsed = []
    for row_number, row in enumerate(values[1:], start=2):
        padded = list(row) + [""] * (len(expected_headers) - len(row))
        parsed.append({header: padded[index] for index, header in enumerate(expected_headers)})
    return parsed


def write_workbook(headers, rows, file_format="xlsx", instructions=None, sheet_name="Datos"):
    if file_format not in ("xlsx", "xls"):
        raise ValueError("Formato de exportación no permitido.")
    if file_format == "xlsx":
        openpyxl = _load_openpyxl()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        # Imported text must remain text, including values beginning with '='.
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        _style_openpyxl(sheet, len(headers))
        if instructions:
            guide = workbook.create_sheet("Instrucciones")
            guide.append(["Campo", "Detalle"])
            for item in instructions:
                guide.append(list(item))
            _style_openpyxl(guide, 2)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    try:
        import xlwt
    except ImportError as error:
        raise ExcelDependencyError("La exportación .xls requiere instalar xlwt.") from error
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet(sheet_name)
    header_style = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour blue; font: colour white;")
    for col, header in enumerate(headers):
        sheet.write(0, col, header, header_style)
        sheet.col(col).width = min(max(len(header) + 4, 14), 45) * 256
    for row_index, row in enumerate(rows, start=1):
        for col, header in enumerate(headers):
            sheet.write(row_index, col, row.get(header, ""))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), "application/vnd.ms-excel"


def write_error_report(errors, file_format="xlsx"):
    headers = ["Fila", "Columna", "Valor", "Motivo"]
    rows = [{header: error.get(header.lower(), "") for header in headers} for error in errors]
    return write_workbook(headers, rows, file_format=file_format)


def parse_bool(value, default=True):
    normalized = str(value or "").strip().lower()
    if not normalized:
        return default
    if normalized in {"si", "sí", "true", "1", "activo", "activa", "yes"}:
        return True
    if normalized in {"no", "false", "0", "inactivo", "inactiva"}:
        return False
    raise ValueError("Usa SI/NO, TRUE/FALSE o 1/0.")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value):
    return " ".join(str(value or "").replace("\ufeff", "").strip().lower().split())


def _style_openpyxl(sheet, column_count):
    from openpyxl.styles import Font, PatternFill

    header = sheet[1]
    for cell in header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F81BD")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [max(14, min(42, len(str(sheet.cell(1, col).value or "")) + 10)) for col in range(1, column_count + 1)]
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + col) if col <= 26 else f"A{col}"].width = width
